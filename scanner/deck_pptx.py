"""
PowerPoint renderer for stakeholder briefs (python-pptx). Renders the same
content model as scanner/deck_pdf.py, so slides and PDFs always match.

Sections are packed onto 16:9 slides: a section that fits joins the previous
slide; a long one gets its own slide, and tables longer than a slide are
continued on the next one.
"""

from __future__ import annotations

import io
import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from scanner.deck_content import Brief, Section, TONES

DARK = RGBColor(11, 22, 40)
LIGHT = RGBColor(244, 246, 251)
TEXT = RGBColor(30, 40, 58)
DIM = RGBColor(110, 125, 148)
WHITE = RGBColor(255, 255, 255)
FONT = "Calibri"

SLIDE_W, SLIDE_H = 13.333, 7.5
MX = 0.6                    # side margin (in)
TOP = 1.45                  # content top (below title band)
BOTTOM = 6.95               # content bottom (above footer)
CW = SLIDE_W - 2 * MX


def rgb(t):
    return RGBColor(*t)


def tint(t, a=0.12):
    return RGBColor(*(int(c * a + 250 * (1 - a)) for c in t))


# ── Height estimation (inches) ───────────────────────────────────────────────
def _lines(text: str, width_in: float, pt: float) -> int:
    chars_per_line = max(10, int(width_in * 72 / (pt * 0.44)))
    return sum(max(1, math.ceil(len(par) / chars_per_line)) for par in str(text).split("\n"))


def _lh(pt: float) -> float:
    return pt * 1.18 / 72


def _row_h(row, col_w, pt=11) -> float:
    return max(_lines(c, w - 0.15, pt) for c, w in zip(row, col_w)) * _lh(pt) + 0.12


def block_height(kind, payload) -> float:
    if kind == "text":
        return _lines(payload, CW, 15) * _lh(15) + 0.15
    if kind == "bullets":
        return sum(_lines(b, CW - 0.35, 15) * _lh(15) + 0.08 for b in payload) + 0.1
    if kind == "kv":
        return sum(_lines(v, CW - 2.6, 14) * _lh(14) + 0.1 for _, v in payload) + 0.1
    if kind == "stats":
        return 1.45
    if kind == "callout":
        return _lines(payload, CW - 0.6, 15) * _lh(15) + 0.5
    if kind == "flow":
        return len(payload) * 0.5 + (len(payload) - 1) * 0.18 + 0.15
    if kind == "table":
        cw = _col_widths(payload)
        return _row_h(payload["cols"], cw) + sum(_row_h(r, cw) for r in payload["rows"]) + 0.15
    return 0.5


def _col_widths(spec):
    tot = sum(spec["widths"])
    return [CW * w / tot for w in spec["widths"]]


# ── Deck builder ─────────────────────────────────────────────────────────────
class DeckWriter:
    def __init__(self, data: dict):
        self.data = data
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]

    # primitives
    def _rect(self, slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = color
        s.line.fill.background()
        s.shadow.inherit = False
        return s

    def _text(self, slide, x, y, w, h, text, pt=14, color=TEXT, bold=False,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = str(text)
        f = r.font
        f.size, f.bold, f.name = Pt(pt), bold, FONT
        f.color.rgb = color
        return tb

    def _notes(self, slide, text):
        slide.notes_slide.notes_text_frame.text = text

    # slides
    def title_slide(self, title, subtitle, accent, sub_pt=24):
        s = self.prs.slides.add_slide(self.blank)
        self._rect(s, 0, 0, SLIDE_W, SLIDE_H, DARK)
        self._rect(s, 0, 0, 0.18, SLIDE_H, rgb(accent))
        self._text(s, 0.9, 0.7, 10, 0.4, "FRESHPAINT  |  PRIVACY-COMPLIANT HEALTHCARE ANALYTICS",
                   pt=12, color=rgb(accent), bold=True)
        self._text(s, 0.9, 2.3, 11.5, 1.6, title, pt=44, color=WHITE, bold=True)
        sub_y = 3.9 if sub_pt <= 28 else 3.2
        self._text(s, 0.9, sub_y, 11.5, 1.2, subtitle, pt=sub_pt, color=rgb(accent), bold=sub_pt > 28)
        self._text(s, 0.9, 6.3, 11.5, 0.4,
                   f"Prepared for {self.data.get('company', '')}  |  {self.data.get('rep_name', '')}  |  "
                   f"{self.data.get('date_str', '')}", pt=12, color=DIM)
        return s

    def persona_cover(self, brief: Brief):
        p = brief.persona
        s = self.title_slide(p["headline"][0], p["headline"][1], p["accent"], sub_pt=44)
        self._text(s, 0.9, 1.5, 11, 0.5, f"FOR THE {p['short']}  -  {p['focus']}",
                   pt=14, color=DIM, bold=True)
        return s

    def content_slide(self, brief: Brief, title: str):
        p = brief.persona
        accent = rgb(p["accent"])
        s = self.prs.slides.add_slide(self.blank)
        self._rect(s, 0, 0, SLIDE_W, 1.1, DARK)
        self._rect(s, 0, 0, 0.12, 1.1, accent)
        self._text(s, MX, 0.14, 8, 0.3, f"FOR THE {p['short']}", pt=10, color=accent, bold=True)
        self._text(s, MX, 0.42, CW, 0.6, title, pt=26, color=WHITE, bold=True)
        # footer
        self._text(s, MX, 7.05, 6, 0.3, f"Freshpaint  |  Prepared for {self.data.get('company', '')}",
                   pt=9, color=DIM)
        self._text(s, SLIDE_W - MX - 6, 7.05, 6, 0.3,
                   "Directional benchmarks and regulatory summary. Not legal advice.",
                   pt=9, color=DIM, align=PP_ALIGN.RIGHT)
        return s

    def subhead(self, slide, brief, y, title) -> float:
        self._rect(slide, MX, y + 0.06, 0.08, 0.28, rgb(brief.persona["accent"]))
        self._text(slide, MX + 0.18, y, CW, 0.4, title.upper(), pt=14, color=TEXT, bold=True)
        return y + 0.48

    # blocks — each returns the new y
    def b_text(self, s, brief, y, t):
        h = block_height("text", t)
        self._text(s, MX, y, CW, h, t, pt=15)
        return y + h

    def b_bullets(self, s, brief, y, items):
        accent = rgb(brief.persona["accent"])
        for it in items:
            h = _lines(it, CW - 0.35, 15) * _lh(15) + 0.08
            self._rect(s, MX + 0.05, y + 0.1, 0.09, 0.09, accent)
            self._text(s, MX + 0.3, y, CW - 0.35, h, it, pt=15)
            y += h
        return y + 0.1

    def b_kv(self, s, brief, y, pairs):
        accent = rgb(brief.persona["accent"])
        for k, v in pairs:
            h = _lines(v, CW - 2.6, 14) * _lh(14) + 0.1
            self._text(s, MX, y, 2.4, h, k, pt=14, color=accent, bold=True)
            self._text(s, MX + 2.5, y, CW - 2.6, h, v, pt=14)
            y += h
        return y + 0.1

    def b_stats(self, s, brief, y, items):
        n, gap, h = len(items), 0.2, 1.3
        w = (CW - gap * (n - 1)) / n
        for i, (val, label, tone) in enumerate(items):
            color = rgb(TONES.get(tone) or brief.persona["accent"])
            x = MX + i * (w + gap)
            self._rect(s, x, y, w, h, LIGHT)
            self._rect(s, x, y, w, 0.07, color)
            pt = 30 if len(val) <= 8 else 24 if len(val) <= 11 else 20
            self._text(s, x, y + 0.15, w, 0.6, val, pt=pt, color=color, bold=True,
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            self._text(s, x + 0.1, y + 0.78, w - 0.2, 0.45, label, pt=11, color=DIM,
                       align=PP_ALIGN.CENTER)
        return y + h + 0.15

    def b_callout(self, s, brief, y, t):
        h = block_height("callout", t) - 0.1
        self._rect(s, MX, y, CW, h, tint(brief.persona["accent"]))
        self._rect(s, MX, y, 0.08, h, rgb(brief.persona["accent"]))
        self._text(s, MX + 0.3, y, CW - 0.5, h, t, pt=15, anchor=MSO_ANCHOR.MIDDLE)
        return y + h + 0.1

    def b_flow(self, s, brief, y, steps):
        accent = rgb(brief.persona["accent"])
        bw = CW * 0.7
        x = MX + (CW - bw) / 2
        for i, st in enumerate(steps):
            hi = i in (1, 2)
            self._rect(s, x, y, bw, 0.5, accent if hi else LIGHT)
            self._text(s, x, y, bw, 0.5, st, pt=13, color=WHITE if hi else TEXT, bold=hi,
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            y += 0.5
            if i < len(steps) - 1:
                self._text(s, x, y - 0.05, bw, 0.25, "↓", pt=12, color=DIM, align=PP_ALIGN.CENTER)
                y += 0.18
        return y + 0.15

    def b_table(self, s, brief, y, spec):
        cw = _col_widths(spec)
        rows = [spec["cols"]] + spec["rows"]
        h = sum(_row_h(r, cw) for r in rows)
        shape = s.shapes.add_table(len(rows), len(cw), Inches(MX), Inches(y), Inches(CW), Inches(h))
        tbl = shape.table
        tbl.first_row = True
        for j, w in enumerate(cw):
            tbl.columns[j].width = Inches(w)
        stripe = tint(brief.persona["accent"], 0.08)
        for i, r in enumerate(rows):
            tbl.rows[i].height = Inches(_row_h(r, cw))
            for j, val in enumerate(r):
                cell = tbl.cell(i, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = DARK if i == 0 else (stripe if i % 2 == 0 else WHITE)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.04)
                tf = cell.text_frame
                tf.word_wrap = True
                tf.paragraphs[0].text = ""
                run = tf.paragraphs[0].add_run()
                run.text = str(val)
                run.font.size = Pt(11)
                run.font.name = FONT
                run.font.bold = i == 0
                run.font.color.rgb = WHITE if i == 0 else TEXT
        return y + h + 0.15

    # packing
    def add_brief(self, brief: Brief, cover: bool = True):
        if cover:
            self.persona_cover(brief)
        avail = BOTTOM - TOP
        slide, y = None, BOTTOM
        for sec in brief.sections:
            for part in _split_section(sec, avail):
                need = sum(block_height(k, p) for k, p in part.blocks)
                if slide is not None and y + 0.5 + need <= BOTTOM:
                    y = self.subhead(slide, brief, y + 0.1, part.title)
                else:
                    slide, y = self.content_slide(brief, part.title), TOP
                for kind, payload in part.blocks:
                    y = getattr(self, "b_" + kind)(slide, brief, y, payload)

    def save(self) -> bytes:
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


def _split_section(sec: Section, avail: float) -> list:
    """Split a section whose blocks overflow one slide, breaking long tables by rows."""
    parts, cur, used = [], [], 0.0
    for kind, payload in sec.blocks:
        h = block_height(kind, payload)
        if kind == "table" and used + h > avail:
            cw = _col_widths(payload)
            head = _row_h(payload["cols"], cw)
            rows, acc = [], used + head + 0.15
            for r in payload["rows"]:
                rh = _row_h(r, cw)
                if acc + rh > avail and rows:
                    cur.append(("table", {**payload, "rows": rows}))
                    parts.append(cur)
                    cur, rows, acc = [], [], head + 0.15
                rows.append(r)
                acc += rh
            cur.append(("table", {**payload, "rows": rows}))
            used = acc
            continue
        if used + h > avail and cur:
            parts.append(cur)
            cur, used = [], 0.0
        cur.append((kind, payload))
        used += h
    if cur:
        parts.append(cur)
    out = []
    for i, blocks in enumerate(parts):
        out.append(Section(sec.title if i == 0 else f"{sec.title} (cont.)", blocks))
    return out


def render_pptx(briefs: list[Brief], data: dict) -> bytes:
    """One deck. A single brief gets its own cover; several get a deck title slide plus a cover each."""
    w = DeckWriter(data)
    if len(briefs) > 1:
        names = ", ".join(b.persona["short"] for b in briefs)
        s = w.title_slide("Stakeholder Briefing", f"{data.get('company', '')}", (0, 200, 150))
        w._notes(s, f"Sections in this deck: {names}.")
    for b in briefs:
        w.add_brief(b, cover=True)
    return w.save()
